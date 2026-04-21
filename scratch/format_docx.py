import docx
from docx.shared import Pt, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import os

def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.first_child_found_in("w:tcBorders")
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders')
        tcPr.append(tcBorders)
    for edge in ('start', 'top', 'end', 'bottom', 'insideH', 'insideV'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = 'w:{}'.format(edge)
            element = tcBorders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                tcBorders.append(element)
            for key in ["sz", "val", "color", "space", "shadow"]:
                if key in edge_data:
                    element.set(qn('w:{}'.format(key)), str(edge_data[key]))

def remove_borders(table):
    for row in table.rows:
        for cell in row.cells:
            set_cell_border(
                cell,
                top={"val": "nil"},
                bottom={"val": "nil"},
                start={"val": "nil"},
                end={"val": "nil"},
                insideH={"val": "nil"},
                insideV={"val": "nil"}
            )

def set_apa_table_borders(table):
    # Remove all borders first
    remove_borders(table)
    # Apply top and bottom border to the first row (Header)
    for cell in table.rows[0].cells:
        set_cell_border(
            cell,
            top={"sz": 12, "val": "single", "color": "000000", "space": "0"},
            bottom={"sz": 12, "val": "single", "color": "000000", "space": "0"}
        )
    # Apply bottom border to the last row
    for cell in table.rows[-1].cells:
        set_cell_border(
            cell,
            bottom={"sz": 12, "val": "single", "color": "000000", "space": "0"}
        )

def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text)
    p.style = doc.styles['Heading 1' if level == 1 else 'Heading 2']
    return p

def main():
    doc = docx.Document()
    
    # Configure default paragraph formatting (APA 7)
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    style.paragraph_format.line_spacing = 2.0
    style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    
    # Configure heading styles
    h1 = doc.styles['Heading 1']
    h1.base_style = style
    h1.font.name = 'Times New Roman'
    h1.font.size = Pt(12)
    h1.font.bold = True
    h1.font.color.rgb = docx.shared.RGBColor(0, 0, 0)
    h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h1.paragraph_format.line_spacing = 2.0
    
    h2 = doc.styles['Heading 2']
    h2.base_style = style
    h2.font.name = 'Times New Roman'
    h2.font.size = Pt(12)
    h2.font.bold = True
    h2.font.color.rgb = docx.shared.RGBColor(0, 0, 0)
    h2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h2.paragraph_format.line_spacing = 2.0
    
    # Layout margins
    for section in doc.sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(2.54)
        section.right_margin = Cm(2.54)
    
    # Title Page
    for _ in range(3):
        doc.add_paragraph()
    
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Implementación de un Software Informático Móvil para Optimizar el Control de Pedidos en el Servicio de Fabricación de Equipos y Maquinarias en la Empresa FAMINSA PERÚ SAC en el Año 2022")
    run.bold = True
    
    doc.add_paragraph() # Blank line after title
    
    p = doc.add_paragraph("Dennys Fabian Falcón Lainez (u20217428) y Shantall Echevarria Soto (u18101900)")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    p = doc.add_paragraph("Ingeniería de Sistemas e Informática, Universidad Tecnológica del Perú")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    p = doc.add_paragraph("Vicky Leonor Alata Linares")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    p = doc.add_paragraph("17 de abril de 2026")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_page_break()
    
    # Abstract Page
    add_heading(doc, "Resumen", level=1)
    
    p_abstract = doc.add_paragraph("La empresa FAMINSA PERÚ SAC fabrica y comercializa máquinas industriales, además de brindar servicios de mantenimiento. Actualmente, presenta deficiencias logísticas al carecer de una herramienta informática adecuada que le permita registrar y controlar la información de los pedidos, generando retrasos. La raíz de este incidente es la carencia de conocimiento técnico y personal especializado en Tecnologías de la Información (TI). Por tanto, este trabajo de investigación busca determinar de qué manera influye implementar un software informático móvil para optimizar el control de pedidos en este servicio. La población abarca 15 trabajadores del área (población universo), y se evalúa a través de recopilación de requerimientos, diseño de interfaces, implementación con metodología ágil Scrum y pruebas de evaluación cuantitativas.")
    p_abstract.paragraph_format.first_line_indent = Inches(0.5)
    
    p_keys = doc.add_paragraph()
    p_keys.paragraph_format.first_line_indent = Inches(0.5)
    run_keys = p_keys.add_run("Palabras clave: ")
    run_keys.italic = True
    p_keys.add_run("software móvil, control de pedidos, mantenimiento de equipos, metodología ágil Scrum, automatización de información.")
    
    doc.add_page_break()

    # Introduction / Body
    title_body = doc.add_paragraph()
    title_body.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_body = title_body.add_run("Implementación de un Software Informático Móvil para Optimizar el Control de Pedidos en el Servicio de Fabricación de Equipos y Maquinarias en la Empresa FAMINSA PERÚ SAC en el Año 2022")
    run_body.bold = True
    
    add_heading(doc, "Idea de la Investigación", level=1)
    
    add_heading(doc, "Tema de Investigación", level=2)
    p = doc.add_paragraph("Diseño de un software informático móvil para optimizar el control de pedidos en el área de servicio de fabricación y mantenimiento de equipos y maquinarias en la empresa FAMINSA PERÚ SAC.")
    p.paragraph_format.first_line_indent = Inches(0.5)
    
    add_heading(doc, "Problemática", level=2)
    p = doc.add_paragraph("La empresa FAMINSA PERÚ SAC se desenvuelve en el sector construcción fabricando y comercializando máquinas industriales, además de brindar servicios de mantenimiento. Sin embargo, en la actualidad presenta deficiencias logísticas, ya que no cuenta con la herramienta informática adecuada que le permita registrar, controlar o almacenar la información de los pedidos de manera íntegra, lo que genera retrasos y un control ineficiente. La raíz de este incidente es la carencia de conocimiento técnico y personal especializado en el área de Tecnologías de la Información (TI), limitando el desempeño de los operarios y el cumplimiento en la entrega oportuna de los pedidos.")
    p.paragraph_format.first_line_indent = Inches(0.5)
    
    add_heading(doc, "Población y Lugar de Estudio", level=2)
    p = doc.add_paragraph("La población está conformada por 15 trabajadores del área de mantenimiento de la empresa Faminsa. Dado que la población es pequeña, se toma el universo y no se requiere el cálculo de una muestra probabilística (Hernández et al., 2014).")
    p.paragraph_format.first_line_indent = Inches(0.5)
    
    add_heading(doc, "Línea de Investigación", level=1)
    p = doc.add_paragraph("La propuesta responde a la línea de investigación de Ingeniería de Sistemas e Informática de la UTP.")
    p.paragraph_format.first_line_indent = Inches(0.5)
    
    add_heading(doc, "Competencias del Modelo del Egresado Desarrolladas", level=1)
    
    table1 = doc.add_table(rows=5, cols=2)
    table1.style = 'Table Grid'
    hdr_cells = table1.rows[0].cells
    hdr_cells[0].text = 'Nro.'
    hdr_cells[1].text = 'Competencias'
    
    row1 = table1.rows[1].cells
    row1[0].text = '1'
    row1[1].text = 'Resolución de problemas de ingeniería: Identificando y evaluando un problema real de ineficiencia de datos logísticos, proporcionando una solución móvil integral.'
    row2 = table1.rows[2].cells
    row2[0].text = '2'
    row2[1].text = 'Análisis y diseño de sistemas informáticos: Aplicación de metodologías ágiles de desarrollo (Scrum / Mobile-D) para crear el modelado, las funcionalidades y la arquitectura del software.'
    row3 = table1.rows[3].cells
    row3[0].text = '3'
    row3[1].text = 'Uso de herramientas modernas: Utilización de recursos de TI y gestores de bases de datos para programar y alojar la infraestructura tecnológica del aplicativo móvil.'
    row4 = table1.rows[4].cells
    row4[0].text = '4'
    row4[1].text = 'Comunicación y trabajo multidisciplinario: Gestión de requerimientos e interacción continua con los miembros y operarios de la organización de FAMINSA PERÚ SAC.'
    
    set_apa_table_borders(table1)
    
    doc.add_paragraph() # Spacing after table
    
    add_heading(doc, "Actividades Clave del Alumno", level=1)
    p = doc.add_paragraph("Las actividades clave que se deben realizar durante el desarrollo de la investigación incluyen:")
    p.paragraph_format.first_line_indent = Inches(0.5)
    
    p = doc.add_paragraph("1. Fase de recopilación y pre-test: Levantamiento de requerimientos y toma de mediciones iniciales mediante cuestionarios (fichas de registro) sobre el cumplimiento en entrega de pedidos y encuestas sobre el desempeño de los operarios.")
    p.paragraph_format.first_line_indent = Inches(0.5)
    p = doc.add_paragraph("2. Fase de diseño de la aplicación: Diseño de interfaces (Balsamiq Mockup), creación de diccionarios de datos e historias de usuario (Product Backlog) del aplicativo móvil en conjunto con el entorno funcional.")
    p.paragraph_format.first_line_indent = Inches(0.5)
    p = doc.add_paragraph("3. Fase de desarrollo e implementación: Programación y automatización de la información del software siguiendo las ceremonias y lineamientos previstos en la metodología ágil Scrum (Sprints).")
    p.paragraph_format.first_line_indent = Inches(0.5)
    p = doc.add_paragraph("4. Fase de evaluación y análisis post-test: Testeo del software con usuarios y medición cuantitativa de las mejoras obtenidas (Post-CEP, Post-MDO), analizadas y contrastadas mediante el entorno SPSS y métodos estadísticos descriptivos o de prueba binomial.")
    p.paragraph_format.first_line_indent = Inches(0.5)
    p = doc.add_paragraph("5. Redacción de informe: Discusión de los resultados experimentales en conformidad con las regulaciones de la universidad y las normas APA 7ma edición.")
    p.paragraph_format.first_line_indent = Inches(0.5)
    
    add_heading(doc, "Autorización y Ética en la Investigación", level=1)
    p = doc.add_paragraph("Para el trabajo de investigación se consultará información de la empresa FAMINSA PERÚ S.A.C. Dicha información no es de dominio público, motivo por el que se requerirá la autorización formal para su empleo y divulgación.")
    p.paragraph_format.first_line_indent = Inches(0.5)
    
    p = doc.add_paragraph("Asimismo, se recomienda que el proyecto resultante sea evaluado y avalado por el comité de ética. Esto se debe a que la investigación manipula datos reales de registro de equipos operados por trabajadores identificables y aplica encuestas a la fuerza laboral de un negocio privado, requiriendo el uso y manipulación cautelosa, así como la conformidad explícita de los participantes.")
    p.paragraph_format.first_line_indent = Inches(0.5)
    
    doc.add_page_break()
    
    add_heading(doc, "Referencias", level=1)
    
    # Reference entry
    p_ref = doc.add_paragraph("Hernández Sampieri, R., Fernández Collado, C., & Baptista Lucio, P. (2014). Metodología de la investigación (6ta ed.). McGraw-Hill Education.")
    p_ref.paragraph_format.left_indent = Inches(0.5)
    p_ref.paragraph_format.first_line_indent = Inches(-0.5)
    
    p_ref2 = doc.add_paragraph("American Psychological Association. (2020). Publication manual of the American Psychological Association (7ma ed.). https://doi.org/10.1037/0000165-000")
    p_ref2.paragraph_format.left_indent = Inches(0.5)
    p_ref2.paragraph_format.first_line_indent = Inches(-0.5)
    
    outpath = r"C:\Users\adhar\Downloads\(AC-S04-ATI1) Avance de trabajo de investigación 1 - Formato APA 7.docx"
    doc.save(outpath)
    print(f"File saved successfully to {outpath}")

if __name__ == '__main__':
    main()
