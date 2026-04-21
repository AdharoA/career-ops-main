import docx

doc_path = r"C:\Users\adhar\Downloads\(AC-S04-ATI1) Avance de trabajo de investigación 1 - Formato APA 7.docx"
out_path = r"C:\Users\adhar\Downloads\(AC-S04-ATI1) Avance de trabajo de investigación 1 - Modificado.docx"

doc = docx.Document(doc_path)

# Find the location of "Referencias" to insert the Intro right before it.
ref_index = -1
for i, p in enumerate(doc.paragraphs):
    if p.text.strip().lower() == "referencias":
        ref_index = i
        break

# We will create the paragraphs to insert
new_paragraphs = [
    ("Introducción de la Revisión Sistemática de Literatura", "Heading 1"),
    ("En la actualidad, el sector de manufactura y servicios de mantenimiento enfrenta grandes desafíos logísticos relacionados con el registro y seguimiento de pedidos. A nivel global, la adopción de la Industria 4.0 y el uso de Tecnologías de la Información (TI), específicamente aplicativos móviles y sistemas de gestión integrados, se han convertido en pilares fundamentales para incrementar la productividad empresarial (Gestión, 2017; JVS Informática, 2020).", "Normal"),
    ("Bajo este contexto específico, la empresa FAMINSA PERÚ SAC, dedicada a la fabricación y comercialización de máquinas industriales, evidencia serias deficiencias logísticas en su área de mantenimiento. La situación actual revela un manejo ineficiente de los requerimientos debido a la ausencia de una herramienta informática adecuada que registre, controle y almacene la información de manera íntegra. Esto contrasta sustancialmente con la situación deseada: un entorno de trabajo automatizado, capaz de garantizar un control eficiente y eliminar los retrasos operativos originados por la recolección manual de información y la carencia de especialización técnica en TI en el rubro.", "Normal"),
    ("El presente trabajo justifica su importancia técnica, económica y social debido a que proporcionar una base tecnológica moderna no solo incrementa el rendimiento laboral del equipo humano de mantenimiento, sino que también resuelve un problema estructural operativo. La implementación de innovaciones usando marcos de trabajo ágiles permite optimizar los cuellos de botella organizacionales, lo que se traduce directamente en un manejo rentable de los recursos.", "Normal"),
    ("En este sentido, resulta altamente relevante llevar a cabo una Revisión Sistemática de Literatura (RSL) enfocada en el impacto de la implementación de software móvil empleando metodologías ágiles en áreas logísticas y de mantenimiento. Una prospección inicial demuestra la existencia de literatura científica pertinente para soportar esta revisión. Por ejemplo, Diego et al. (2020) diseñaron un aplicativo web y móvil con metodología Scrum para agilizar la gestión de vehículos, demostrando un aumento cuantitativo en la eficiencia de los tiempos de los procesos. Paralelamente, Enrique et al. (2020) estandarizaron necesidades operativas de una casa editora aplicando la metodología iterativa Scrum para desplegar una plataforma móvil gratuita y escalable. De manera similar orientada a pedidos, Jaime et al. (2021) desarrollaron una aplicación empleando el método iterativo XP, logrando incrementar los niveles de servicio y respuestas a clientes de una clínica veterinaria. Asimismo, Antonio et al. (2021) evidenciaron que el uso de sistemas móviles (Mobile-D) para gestión documental y control logró minimizar notablemente el tiempo promedio de ubicación de datos desfasados. Además, Quezada et al. (2017) corroboraron que mecanizar recursos operativos mediante tecnologías móviles optimiza exponencialmente los medios económicos y humanos disponibles.", "Normal"),
    ("A partir de estas cinco investigaciones recientes (2017-2021), se confirma un consenso empírico que valida el uso conjunto de metodologías iterativas y tecnología móvil en la optimización de procesos operativos. Sin embargo, no se ha identificado previamente una RSL consolidada en la literatura que organice sistemáticamente los principales retos, métricas de éxito y variables al aplicar frameworks específicos como Scrum con miras al sector del mantenimiento de maquinaria industrial pesada; vacío fundamental que este trabajo propone abordar.", "Normal")
]

# Insert the paragraphs BEFORE "Referencias"
if ref_index != -1:
    ref_paragraph = doc.paragraphs[ref_index]
    
    for text, style in new_paragraphs:
        new_p = ref_paragraph.insert_paragraph_before(text)
        new_p.style = doc.styles[style]
else:
    # Append if not found
    for text, style in new_paragraphs:
        doc.add_paragraph(text, style=docx.enum.style.WD_STYLE_TYPE.PARAGRAPH)

# Also, append references at the end:
doc.add_paragraph("Antonio et al. (2021). Aplicativo móvil para mejorar la gestión documental de vehículos. Asociación de Huamachuco.", style="Normal")
doc.add_paragraph("Diego et al. (2020). Ejecución de un sistema web y móvil para una gestión ágil del control de vehículos. Universidad y Ciencia.", style="Normal")
doc.add_paragraph("Enrique et al. (2020). Creación y aplicación de un software en la casa editora Uleam. Repositorio Uleam.", style="Normal")
doc.add_paragraph("Jaime et al. (2021). Construcción de un aplicativo móvil con el fin de mejorar y gestionar pedidos. Clínica Janavet.", style="Normal")
doc.add_paragraph("Quezada, P. et al. (2017). Creación de un sistema web y móvil para gestionar los vehículos basados en características y metodologías ágiles. UTPL.", style="Normal")

doc.save(out_path)
print(f"File successfully modified and saved to: {out_path}")
