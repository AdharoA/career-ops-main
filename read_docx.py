import docx
import sys

def extract_docx_content(filepath):
    try:
        doc = docx.Document(filepath)
        print("=== PARAGRAPHS ===")
        for i, p in enumerate(doc.paragraphs):
            if p.text.strip():
                print(f"[{i}] {p.style.name}: {p.text}")
        
        print("\n=== TABLES ===")
        for i, table in enumerate(doc.tables):
            print(f"Table {i+1}:")
            for row in table.rows:
                row_data = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
                print(" | ".join(row_data))
            print("-" * 20)
    except Exception as e:
        print(f"Error: {e}")

if __name__ == '__main__':
    extract_docx_content(sys.argv[1])
